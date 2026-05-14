"""Document loader — reads .docx and .pdf, attaches minimal metadata."""
from __future__ import annotations
import hashlib, re
from pathlib import Path

import docx
import fitz
from llama_index.core import Document


_UA = re.compile(r"UA-(\d+)\s*(.*)")


def _ua_meta(path: Path) -> dict:
    d = path.parent.name
    m = _UA.match(d)
    return {
        "ua_id":    f"UA-{m.group(1)}" if m else "UA-?",
        "ua_title": (m.group(2).strip() if m else d),
    }


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _read_docx(p: Path) -> str:
    d = docx.Document(str(p))
    parts = [x.text for x in d.paragraphs if x.text.strip()]
    for tbl in d.tables:
        for row in tbl.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _read_pdf(p: Path) -> str:
    with fitz.open(p) as doc:
        return "\n\n".join(f"[page {i + 1}]\n{pg.get_text('text')}"
                           for i, pg in enumerate(doc))


def load_corpus(root: Path) -> list[Document]:
    docs: list[Document] = []
    for p in root.rglob("*"):
        if p.suffix.lower() == ".docx":
            text = _read_docx(p)
        elif p.suffix.lower() == ".pdf":
            text = _read_pdf(p)
        else:
            continue
        if not text.strip():
            continue
        docs.append(Document(
            text=text,
            metadata={
                **_ua_meta(p),
                "source_path": str(p),
                "file_name": p.name,
                "doc_type": p.suffix.lower().lstrip("."),
                "sha256": _sha256(p),
            },
        ))
    return docs